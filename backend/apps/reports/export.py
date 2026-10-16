"""File generation (B7): CSV (UTF-8 BOM), XLSX (openpyxl write-only), PDF (reportlab, A4 landscape)."""
import csv
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

CONTENT_TYPES = {"CSV": "text/csv; charset=utf-8",
                 "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                 "PDF": "application/pdf"}
EXTENSIONS = {"CSV": "csv", "XLSX": "xlsx", "PDF": "pdf"}
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_text(v):
    """Text that could be read as a formula gets a leading apostrophe. Numbers are left alone."""
    if isinstance(v, str) and v and v[0] in _FORMULA_START:
        return "'" + v
    return v


def _plain(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    return v


def write_csv(path, columns, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow([c.label for c in columns])
        for r in rows:
            w.writerow([safe_text(_plain(v)) for v in r])


def write_xlsx(path, title, columns, rows, summary=None):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=(title or "Report")[:31])
    for i, c in enumerate(columns, 1):
        ws.column_dimensions[_col(i)].width = min(max(len(c.label) + 2, 12), 40)
    head = []
    for c in columns:
        cell = WriteOnlyCell(ws, value=c.label)
        cell.font = Font(bold=True)
        head.append(cell)
    ws.append(head)
    for r in rows:
        ws.append([_xl(v) for v in r])
    if summary:
        s = wb.create_sheet(title="Summary")
        s.append(["Item", "Value"])
        for k, v in summary.items():
            s.append([str(k), _xl(v)])
    wb.save(path)


def _col(i):
    s = ""
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


def _xl(v):
    v = safe_text(_plain(v))
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def _pdf_text(v):
    v = _plain(v)
    if isinstance(v, float):
        v = f"{v:,.2f}"
    return escape(str(v)[:300])


def write_pdf(path, *, title, facility_name, params_text, generated_by, generated_at, columns, rows, summary=None):
    page = landscape(A4)
    doc = SimpleDocTemplate(str(path), pagesize=page, leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=30,
                            title=title)
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7, leading=8.5)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    story = [Paragraph(escape(title), styles["Title"]),
             Paragraph(escape(f"Facility: {facility_name}"), styles["Normal"]),
             Paragraph(escape(f"Parameters: {params_text or 'none'}"), styles["Normal"]),
             Paragraph(escape(f"Generated: {generated_at} by {generated_by}"), styles["Normal"])]
    if summary:
        story.append(Paragraph(escape("Summary: " + "; ".join(f"{k}: {_plain(v)}" for k, v in summary.items())),
                               styles["Normal"]))
    story.append(Spacer(1, 8))
    data = [[Paragraph(escape(c.label), head) for c in columns]]
    data += [[Paragraph(_pdf_text(v), cell) for v in r] for r in rows]
    width = (page[0] - 48) / max(1, len(columns))
    table = Table(data, colWidths=[width] * len(columns), repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7D007D")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F0F6")])]))
    story.append(table)

    def footer(canvas, d):
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(page[0] - 24, 14, f"Page {d.page}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
